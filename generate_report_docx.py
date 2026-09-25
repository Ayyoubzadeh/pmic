#!/usr/bin/env python3
"""Generate Methods & Results DOCX for the AllStrainsExceptResistant pipeline run."""

from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(".")
PLOTS = ROOT / "plots"
FP_REG = PLOTS / "fp_highlights"
FP_CLS = PLOTS / "fp_highlights_cls"
OUT = ROOT / "pMIC_Pipeline_v3_AllStrains_Methods_Results_Report.docx"


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def H(doc, text, level=1):
    return doc.add_heading(text, level=level)


def P(doc, text, bold=False, italic=False, size=11):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    r.font.name = "Calibri"
    p.paragraph_format.space_after = Pt(8)
    return p


def bullets(doc, items):
    for it in items:
        p = doc.add_paragraph(it, style="List Bullet")
        for r in p.runs:
            r.font.size = Pt(11)
            r.font.name = "Calibri"


def table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = h
        shade(cell, "1F4E79")
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(255, 255, 255)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.rows[ri + 1].cells[ci]
            cell.text = str(val)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
            if ri % 2:
                shade(cell, "F2F2F2")
    doc.add_paragraph()
    return t


def fig(doc, path, caption, width=5.6):
    path = Path(path)
    if not path.exists():
        P(doc, f"[Missing figure: {path}]", italic=True, size=10)
        return
    doc.add_picture(str(path), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    r.italic = True
    r.font.size = Pt(10)
    r.font.name = "Calibri"
    cap.paragraph_format.space_after = Pt(12)


def build():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2.0)
        s.bottom_margin = Cm(2.0)
        s.left_margin = Cm(2.2)
        s.right_margin = Cm(2.2)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("pMIC Prediction Pipeline v3")
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor(31, 78, 121)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(
        "Methods and Results Report\n"
        "Primary dataset: AllStrainsExceptResistant  |  "
        "Two-stage SMILES deduplication (median pMIC)  |  "
        "OOF threshold  |  CV Average Precision  |  Sensitivity analysis"
    )
    r.italic = True
    r.font.size = Pt(11)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(
        "Unique molecules after dedup: n = 18,918  |  "
        "Activity cutoff: pMIC >= 6.0  |  "
        "OOF decision threshold = 0.050  |  Seed = 42  |  25 September 2026"
    )
    r.font.size = Pt(10)

    # ── 1 Overview ──────────────────────────────────────────────────────────
    H(doc, "1. Overview", 1)
    P(doc,
      "This report documents the retrain of the antimycobacterial pMIC QSAR pipeline on "
      "the revised modeling table AllStrainsExceptResistant.xlsx. After two-stage SMILES "
      "deduplication (exact string, then RDKit canonical SMILES) with median pMIC, the "
      "workflow used multi-fingerprint representation, a Butina/Tanimoto cluster split "
      "with GroupKFold, stacking regression, binary classification at pMIC >= 6 selected "
      "by out-of-fold Average Precision, OOF decision-threshold optimization frozen before "
      "test evaluation, signed per-bit SHAP mapping, labeled external validation, "
      "sensitivity analysis of the frozen model on H37Rv and Resistant tables, and "
      "virtual screening of approved drugs, 4FDN natural products, and CyanoMetDB.")

    H(doc, "1.1 Objectives", 2)
    bullets(doc, [
        "Retrain the main model on AllStrainsExceptResistant as the sole modeling set.",
        "Collapse exact-SMILES and canonical-SMILES duplicates to one median pMIC on every labeled table, including External, H37Rv, and Resistant.",
        "Plot the pMIC distribution before the train/test split and assess class balance at cutoff 6.",
        "Select classifiers by GroupKFold Average Precision (active ranking under imbalance).",
        "Optimize the classification decision threshold on OOF probabilities only, then evaluate Test, External, and screens at that frozen threshold.",
        "Run sensitivity analysis of the frozen main model on H37Rv and Resistant, reporting overlap-aware metrics.",
        "Screen approved, 4FDN, and CyanoMetDB libraries with the same frozen models.",
        "Map influential Morgan/MACCS bits individually with signed (positive/negative) SHAP contribution.",
    ])

    # ── 2 Methods ───────────────────────────────────────────────────────────
    H(doc, "2. Methods", 1)

    H(doc, "2.1 Datasets and endpoint", 2)
    P(doc,
      "The primary modeling file was AllStrainsExceptResistant.xlsx (29,067 rows as loaded; "
      "29,065 rows with numeric pMIC). pMIC is -log10(MIC [M]). Two additional labeled "
      "tables, H37Rv.xlsx and Resistant.xlsx, were not used for training; they were scored "
      "with the frozen main model for sensitivity analysis. Cleaned_External_Validation.xlsx "
      "was used only after model selection. Unlabeled screening libraries: approved drugs, "
      "4FDN natural-product docking set, and CyanoMetDB V03 2024. COCONUT was not rescreened "
      "in this run (models were saved so that predict_coconut.py can be applied later).")

    H(doc, "2.2 Two-stage SMILES deduplication", 2)
    P(doc,
      "Before fingerprints were computed, each labeled table was collapsed in two stages. "
      "(1) Exact SMILES: strings were stripped and grouped; when the same SMILES appeared "
      "in multiple assay rows with different pMIC values, a single median pMIC was retained "
      "(other columns: first occurrence). (2) Canonical SMILES: RDKit Chem.MolToSmiles(..., "
      "canonical=True) was applied; invalid structures and zero-heavy-atom molecules were "
      "dropped; remaining canonical duplicates were again aggregated by median pMIC. "
      "Unlabeled tables would keep the first occurrence. Salt stripping was not applied. "
      "Median was chosen over the mean as a robust summary of conflicting assay conditions.")
    P(doc,
      "On AllStrainsExceptResistant this reduced 29,065 usable rows to 18,918 unique molecules "
      "(10,147 exact duplicates removed; 3,014 exact-SMILES groups had conflicting pMIC). "
      "After the exact-string collapse, no additional canonical duplicates remained. "
      "pMIC ranged from 1.87 to 9.50.")

    H(doc, "2.3 Class balance, SMOTE policy, and pre-split plot", 2)
    P(doc,
      "A binned pMIC bar plot and Active/Inactive counts at cutoff 6 were drawn on the "
      "unique modeling set before the Butina split. At pMIC >= 6 there were 2,832 actives "
      "out of 18,918 (15.0%). SMOTE was not applied: interpolating in mixed Morgan/MACCS/"
      "descriptor space yields chemically invalid bits. Class imbalance was handled with "
      "class_weight='balanced' and OOF decision-threshold calibration.")

    H(doc, "2.4 Molecular representation", 2)
    bullets(doc, [
        "Morgan fingerprints (radius 2, 2048 bits)",
        "MACCS keys (167 bits)",
        "Atom-pair fingerprints (2048 bits)",
        "RDKit path fingerprints (2048 bits)",
        "216 RDKit descriptors (Ipc excluded) -> 6,527 features in total",
    ])

    H(doc, "2.5 Feature selection", 2)
    bullets(doc, [
        "VarianceThreshold (0.01)",
        "Pearson correlation filter (|r| > 0.95)",
        "Mutual-information SelectKBest: top 1000 (regression) / top 300 (classification)",
        "Selectors fit on training data only and stored in a picklable FeatureTransformer",
    ])

    H(doc, "2.6 Butina cluster split", 2)
    P(doc,
      "Molecules were clustered with Butina on Morgan fingerprints (radius 2, 1024 bits) "
      "at Tanimoto distance cutoff 0.4 (similarity >= 0.6). Entire clusters were assigned "
      "to train or test (~80/20). Within-train CV used GroupKFold by cluster (5 folds).")
    P(doc,
      "Result: 4,368 clusters (2,091 singletons); Train n = 15,134 (2,527 clusters); "
      "Test n = 3,784 (1,841 clusters). Roles were written to preprocessed_data.xlsx.")

    H(doc, "2.7 Models", 2)
    P(doc, "Regression candidates:", bold=True)
    bullets(doc, [
        "RandomForest, ExtraTrees, HistGradientBoosting, Ridge, LinearSVR, XGBoost, LightGBM",
        "StackingRegressor of the top-3 trees by CV R2 -> Ridge",
    ])
    P(doc, "Classification candidates:", bold=True)
    bullets(doc, [
        "Calibrated LinearSVC, LogisticRegression (balanced)",
        "RandomForest, ExtraTrees, HistGradientBoosting (balanced), XGBoost, LightGBM",
        "StackingClassifier of the top-3 trees by CV Average Precision -> LogisticRegression",
        "Gray-zone exclusion: |pMIC - cutoff| < 0.5 removed from classification training only",
    ])

    H(doc, "2.8 Classification cutoff and model-selection metric", 2)
    P(doc,
      "Binary Active labels were compared at pMIC >= 5.5 and >= 6.0 using GroupKFold "
      "random-forest probes (cutoff 7 excluded a priori). Cutoff 6.0 was selected among "
      "near-AUC candidates (soft preference for 6.0). A three-class probe "
      "(Inactive < 5 / Moderate 5-6 / Active >= 6) was reported but did not replace the binary task.")
    P(doc,
      "Classifiers were ranked by GroupKFold out-of-fold Average Precision (PR-AUC). "
      "AP measures ranking of the minority active class under imbalance and matches the "
      "screening goal of recovering actives. ROC-AUC and F1 were reported but were not "
      "used for model selection. Test metrics were not used to pick the model.")

    H(doc, "2.9 OOF decision-threshold optimization", 2)
    P(doc,
      "After the best classifier was chosen by CV AP, the decision threshold on P(active) "
      "was optimized by maximizing balanced accuracy on that model's OOF probabilities "
      "(grid 0.05-0.95). The threshold was then frozen and applied to Train hard labels "
      "for reporting, to Test evaluation and plots, to External validation, to sensitivity "
      "sets, and to library screens. Test labels were not used to choose the threshold.")

    H(doc, "2.10 Applicability domain", 2)
    P(doc,
      "Williams plots used StandardScaler + PCA (up to 50 components) fitted on training "
      "selected features only. Leverage used the hat matrix with h* = 3(p+1)/n_train. "
      "Labeled compounds were in-AD if h <= h* and |standardized residual| <= 3. External "
      "points were overlaid on Train/Test Williams plots. Full library AD-coverage bars "
      "from the previous Smiles.xlsx run were not recomputed in this execution.")

    H(doc, "2.11 Signed fingerprint mapping", 2)
    P(doc,
      "TreeSHAP on the regression LightGBM surrogate ranked Morgan/MACCS bits globally. "
      "For each of the five test molecules with highest predicted pMIC, per-molecule SHAP "
      "was computed and the top bits were drawn individually: green = positive contribution "
      "(increases predicted pMIC), red = negative. The same per-bit signed mapping was "
      "repeated for classification (P_active) on the five test molecules with highest "
      "active probability. Overlay drawings of reference TB drugs used the global top bits.")

    H(doc, "2.12 External validation, sensitivity, and screening", 2)
    P(doc,
      "The frozen Stacking regressor and LightGBM classifier were applied to the "
      "median-deduplicated external set. Canonical SMILES overlap with the modeling set "
      "was reported; metrics were also computed after dropping the overlapping molecules.")
    P(doc,
      "Sensitivity analysis scored H37Rv and Resistant with the same frozen models. "
      "Because H37Rv is nested inside AllStrainsExceptResistant, metrics were split into "
      "all / overlap with modeling / overlap with Train / held-out versus Train. "
      "Resistant has a true held-out subset versus the modeling set.")
    P(doc,
      "Approved, 4FDN, and CyanoMetDB were predicted in batch. Active_predicted used the "
      "frozen OOF threshold (0.050), not 0.5. Invalid SMILES were skipped.")

    H(doc, "2.13 Software", 2)
    bullets(doc, [
        "Python; RDKit; scikit-learn; XGBoost; LightGBM; SHAP; NumPy; Pandas; Matplotlib",
        "Key scripts: pmic_prediction.py, pmic_extras.py, predict_forxlsx.py",
        "Log: pipeline_allstrains.log; cache: cache_features_allstrains.joblib",
    ])

    # ── 3 Results ───────────────────────────────────────────────────────────
    H(doc, "3. Results", 1)

    H(doc, "3.1 Deduplication", 2)
    table(doc,
          ["Set", "Raw rows with pMIC", "Exact dups removed", "Exact conflict pMIC groups",
           "Canonical dups removed", "Unique molecules"],
          [
              ["AllStrainsExceptResistant (main)", "29,065", "10,147", "3,014", "0", "18,918"],
              ["External validation", "122", "30", "3", "0", "92"],
              ["H37Rv (sensitivity)", "21,022", "5,361", "1,959", "0", "15,661"],
              ["Resistant (sensitivity)", "6,798", "4,403", "727", "0", "2,395"],
          ])
    P(doc, "Table 1. Two-stage SMILES collapse with median pMIC.", italic=True, size=10)

    table(doc,
          ["Pair", "Intersection", "Percent of the second set"],
          [
              ["Modeling ∩ External", "2 / 92", "2.2%"],
              ["Modeling ∩ H37Rv", "15,661 / 15,661", "100% (H37Rv is nested)"],
              ["Modeling ∩ Resistant", "2,039 / 2,395", "85.1%"],
          ])
    P(doc, "Table 2. Canonical SMILES overlap with the unique modeling set.", italic=True, size=10)

    H(doc, "3.2 pMIC distribution before the split", 2)
    P(doc,
      "On 18,918 unique modeling molecules, 2,832 (15.0%) had pMIC >= 6 and 16,086 were "
      "inactive at that cutoff. This is moderate-to-strong imbalance and motivated AP-based "
      "model selection and OOF threshold calibration rather than SMOTE.")
    fig(doc, PLOTS / "pmic_distribution.png",
        "Figure 1. pMIC distribution on AllStrainsExceptResistant after median dedup, "
        "before the train/test split (cutoff = 6).")

    H(doc, "3.3 Cutoff comparison (5.5 vs 6.0)", 2)
    table(doc,
          ["Cutoff", "Train active %", "CV ROC-AUC", "CV AP", "CV BalAcc", "CV F1", "CV MCC"],
          [
              ["5.5", "29.5%", "0.826", "0.711", "0.710", "0.592", "0.478"],
              ["6.0 (selected)", "16.2%", "0.863", "0.633", "0.706", "0.547", "0.496"],
          ])
    P(doc, "Table 3. GroupKFold RF probes on the unique-molecule train set. Cutoff 6.0 was selected.",
      italic=True, size=10)
    P(doc,
      "Three-class probe train counts Inactive/Moderate/Active = 7,921 / 4,757 / 2,456; "
      "CV macro-F1 = 0.588; Test macro-F1 = 0.561. Binary classification at pMIC >= 6 "
      "remained the primary task.")

    H(doc, "3.4 Regression (global model, Butina split)", 2)
    P(doc,
      "After variance (4,977 features), correlation (4,922), and mutual-information selection "
      "(top 1,000), Stacking of LightGBM + XGBoost + ExtraTrees -> Ridge was best by CV R2. "
      "Train pMIC mean = 5.03 (SD 0.98; range 2.17-9.50).")
    table(doc,
          ["Set", "R2", "RMSE", "MAE"],
          [
              ["Train", "0.928", "0.263", "0.196"],
              ["CV (OOF)", "0.441", "0.733", "0.561"],
              ["Test", "0.412", "0.710", "0.530"],
          ])
    P(doc, "Table 4. Global regression metrics — Stacking (selected by CV R2).", italic=True, size=10)

    table(doc,
          ["Model", "CV R2", "Test R2", "CV RMSE", "Test RMSE"],
          [
              ["RandomForest", "0.393", "0.364", "0.765", "0.739"],
              ["ExtraTrees", "0.404", "0.385", "0.758", "0.727"],
              ["HistGradBoost", "0.398", "0.351", "0.761", "0.747"],
              ["Ridge", "0.164", "0.086", "0.897", "0.886"],
              ["LinearSVR", "0.137", "0.031", "0.911", "0.912"],
              ["XGBoost", "0.415", "0.368", "0.750", "0.737"],
              ["LightGBM", "0.434", "0.401", "0.738", "0.717"],
              ["Stacking (selected)", "0.441", "0.412", "0.733", "0.710"],
          ])
    P(doc, "Table 5. Regression model comparison on the AllStrains Butina split.", italic=True, size=10)

    for name, cap in [
        ("reg_00_model_comparison.png", "Figure 2. Regression model comparison (CV vs Test R2)."),
        ("reg_01_distribution.png", "Figure 3. Train/Test pMIC and MIC distributions."),
        ("reg_02_pred_vs_exp.png", "Figure 4. Predicted vs experimental pMIC."),
        ("reg_03_residuals.png", "Figure 5. Residual plots."),
        ("reg_04_error_distribution.png", "Figure 6. Error distributions."),
        ("reg_05_metrics_summary.png", "Figure 7. Regression metrics summary."),
        ("reg_06_cv_performance.png", "Figure 8. Per-fold CV regression metrics."),
        ("reg_07_embedding.png", "Figure 9. Chemical-space embeddings (regression)."),
        ("reg_08_shap_summary.png", "Figure 10. Regression SHAP summary."),
        ("reg_09_feature_importance.png", "Figure 11. Regression feature importance."),
        ("reg_10_shap_dependence.png", "Figure 12. Regression SHAP dependence."),
        ("reg_11_y_randomization.png", "Figure 13. Y-randomization (regression)."),
        ("reg_12_williams_plot.png",
         "Figure 14. Williams plot (regression AD) with Train, Test, and External."),
        ("reg_13_rmse_comparison.png", "Figure 15. RMSE comparison across regressors."),
    ]:
        fig(doc, PLOTS / name, cap)

    H(doc, "3.5 Classification (pMIC >= 6.0)", 2)
    P(doc,
      "Gray-zone exclusion removed 3,206 train compounds with |pMIC - 6| < 0.5. "
      "Classification training then used 11,928 compounds (1,263 active, 10.6%); the test "
      "set retained 3,784 (376 active, 9.9%). Feature selection reduced the space to 300 "
      "mutual-information features. LightGBM had the best CV Average Precision (0.751) and "
      "was saved. Stacking CV AP was 0.742. ExtraTrees had a slightly higher Test AUC "
      "(0.814) but was not selected, because selection used CV AP only.")
    table(doc,
          ["Model", "CV AP", "CV AUC", "Test AP", "Test AUC"],
          [
              ["LinearSVM", "0.498", "0.820", "0.285", "0.730"],
              ["LogisticReg", "0.505", "0.822", "0.281", "0.733"],
              ["RandomForest", "0.726", "0.913", "0.439", "0.800"],
              ["ExtraTrees", "0.729", "0.914", "0.454", "0.814"],
              ["HistGradBoost", "0.712", "0.899", "0.395", "0.764"],
              ["XGBoost", "0.734", "0.909", "0.420", "0.779"],
              ["LightGBM (selected)", "0.751", "0.911", "0.435", "0.784"],
              ["Stacking", "0.742", "0.914", "0.454", "0.814"],
          ])
    P(doc, "Table 6. Classification comparison at pMIC >= 6. Best model selected by CV AP.",
      italic=True, size=10)

    P(doc,
      "OOF threshold optimization on LightGBM probabilities selected t = 0.050 (the lower "
      "grid bound). For this class-weighted LightGBM, many inactive probabilities sit near "
      "zero, so t = 0.050 still yielded Test specificity 0.945 and Test recall 0.426. "
      "Ranking metrics (AUC/AP) are the primary quality measures for screening.")
    fig(doc, PLOTS / "cls_00c_oof_threshold.png",
        "Figure 16. OOF decision-threshold optimization (balanced accuracy and F1). "
        "Selected t = 0.050 was frozen before Test evaluation.")

    table(doc,
          ["Set", "Accuracy", "Bal. Acc.", "ROC-AUC", "Avg. Prec.", "F1", "Precision",
           "Recall", "Specificity", "MCC"],
          [
              ["Train", "0.996", "0.998", "1.000", "1.000", "0.982", "0.964", "1.000", "0.996", "0.980"],
              ["CV (OOF)", "0.918", "0.842", "0.911", "0.751", "0.659", "0.590", "0.745", "0.939", "0.618"],
              ["Test", "0.894", "0.685", "0.784", "0.435", "0.443", "0.461", "0.426", "0.945", "0.384"],
          ])
    P(doc, "Table 7. LightGBM metrics at the frozen OOF threshold (0.050).", italic=True, size=10)

    for name, cap in [
        ("cls_00_model_comparison.png", "Figure 17. Classification model comparison (ROC-AUC)."),
        ("cls_00b_model_comparison_ap.png", "Figure 18. Classification model comparison (Average Precision)."),
        ("cls_01_class_distribution.png", "Figure 19. Active/Inactive distributions."),
        ("cls_02_roc_curve.png", "Figure 20. ROC curves."),
        ("cls_03_pr_curve.png", "Figure 21. Precision-Recall curves."),
        ("cls_04_confusion_matrix.png", "Figure 22. Confusion matrices."),
        ("cls_05_metrics_summary.png", "Figure 23. Classification metrics summary."),
        ("cls_06_cv_performance.png", "Figure 24. Per-fold CV classification metrics."),
        ("cls_07_embedding.png", "Figure 25. Chemical-space embeddings (classification)."),
        ("cls_08_shap_summary.png", "Figure 26. Classification SHAP summary."),
        ("cls_09_feature_importance.png", "Figure 27. Classification feature importance."),
        ("cls_10_shap_dependence.png", "Figure 28. Classification SHAP dependence."),
        ("cls_11_y_randomization.png", "Figure 29. Y-randomization (classification)."),
        ("cls_12_f1_comparison.png", "Figure 30. F1 comparison across classifiers."),
        ("cls_13_williams_plot.png",
         "Figure 31. Williams plot (classification AD) with Train, Test, and External."),
    ]:
        fig(doc, PLOTS / name, cap)

    H(doc, "3.6 Signed fingerprint highlights", 2)
    P(doc,
      "Global regression SHAP fingerprint bits were led by maccs_86, maccs_107, maccs_119, "
      "and morgan_715. Per-molecule drawings color atoms green when the bit increases "
      "predicted pMIC and red when it decreases it. Classification SHAP (P_active) was "
      "dominated by morgan_1088 and maccs_98.")
    table(doc,
          ["Rank", "Regression feature", "Mean |SHAP|", "Classification feature", "Mean |SHAP|"],
          [
              ["1", "maccs_86", "0.0061", "morgan_1088", "0.1623"],
              ["2", "maccs_107", "0.0058", "maccs_98", "0.0403"],
              ["3", "maccs_119", "0.0058", "morgan_559", "0.0323"],
              ["4", "morgan_715", "0.0049", "morgan_354", "0.0287"],
              ["5", "morgan_1152", "0.0040", "morgan_716", "0.0210"],
              ["6", "maccs_95", "0.0038", "morgan_115", "0.0205"],
              ["7", "maccs_110", "0.0027", "morgan_317", "0.0164"],
              ["8", "maccs_111", "0.0026", "morgan_522", "0.0150"],
          ])
    P(doc, "Table 8. Top global SHAP fingerprint features (regression vs classification).",
      italic=True, size=10)

    for name, cap in [
        ("test_top1_pMIC_8p42.png",
         "Figure 32. Test top-1 by predicted pMIC — overlay of signed SHAP bits (green +, red -)."),
        ("test_top2_pMIC_8p08.png", "Figure 33. Test top-2 — signed SHAP overlay."),
        ("test_top3_pMIC_7p74.png", "Figure 34. Test top-3 — signed SHAP overlay."),
        ("test_top4_pMIC_7p61.png", "Figure 35. Test top-4 — signed SHAP overlay."),
        ("test_top5_pMIC_7p50.png", "Figure 36. Test top-5 — signed SHAP overlay."),
    ]:
        fig(doc, FP_REG / name, cap, width=5.2)

    P(doc,
      "Individual bits for the top-1 test molecule (all positive on this structure) and one "
      "negative contribution from top-2 (maccs_109) are shown below. The full per-bit catalog "
      "is plots/fp_highlights/per_bit_signed_shap.csv and the test_top*_bits folders.")
    for name, cap in [
        ("test_top1_bits/01_maccs_83_pos.png",
         "Figure 37. Test top-1, bit maccs_83 — positive contribution to predicted pMIC."),
        ("test_top1_bits/03_maccs_119_pos.png",
         "Figure 38. Test top-1, bit maccs_119 — positive contribution."),
        ("test_top1_bits/04_maccs_107_pos.png",
         "Figure 39. Test top-1, bit maccs_107 — positive contribution."),
        ("test_top1_bits/08_morgan_1160_pos.png",
         "Figure 40. Test top-1, bit morgan_1160 — positive contribution."),
        ("test_top2_bits/03_maccs_109_neg.png",
         "Figure 41. Test top-2, bit maccs_109 — negative contribution (decreases predicted pMIC)."),
        ("test_top4_bits/08_morgan_1152_neg.png",
         "Figure 42. Test top-4, bit morgan_1152 — negative contribution."),
    ]:
        fig(doc, FP_REG / name, cap, width=5.0)

    for name, cap in [
        ("map_Isoniazid.png", "Figure 43. Isoniazid — global FP overlay."),
        ("map_Rifampin.png", "Figure 44. Rifampin — global FP overlay."),
        ("map_Bedaquiline.png", "Figure 45. Bedaquiline — global FP overlay."),
        ("map_Delamanid.png", "Figure 46. Delamanid — global FP overlay."),
        ("map_Linezolid.png", "Figure 47. Linezolid — global FP overlay."),
    ]:
        fig(doc, FP_REG / name, cap, width=5.0)

    P(doc, "Classification per-bit examples (P_active SHAP, green increases active probability):")
    for name, cap in [
        ("test_top1_Pactive_1p00.png",
         "Figure 48. Classification overlay on the highest-P(active) test molecule."),
        ("test_top1_bits/01_maccs_98_pos.png",
         "Figure 49. Classification bit maccs_98 — positive contribution to P(active)."),
        ("test_top1_bits/03_morgan_1088_neg.png",
         "Figure 50. Classification bit morgan_1088 — negative contribution to P(active) on this molecule."),
    ]:
        fig(doc, FP_CLS / name, cap, width=5.0)

    H(doc, "3.7 External validation", 2)
    P(doc,
      "On 92 unique external molecules (7 actives at pMIC >= 6) the global regressor reached "
      "R2 = 0.596, RMSE = 0.540, MAE = 0.436. Classification at cutoff 6.0 and frozen "
      "threshold 0.050 gave ROC-AUC = 0.897, AP = 0.303, balanced accuracy = 0.798, "
      "F1 = 0.455, MCC = 0.428. Two external molecules (2.2%) shared canonical SMILES with "
      "the modeling set. Held-out-only metrics on the remaining 90 molecules were essentially "
      "unchanged (R2 = 0.598, AUC = 0.895, AP = 0.303).")
    table(doc,
          ["Subset", "n", "R2", "RMSE", "MAE", "ROC-AUC", "AP", "BalAcc", "F1"],
          [
              ["All unique External", "92", "0.596", "0.540", "0.436", "0.897", "0.303", "0.798", "0.455"],
              ["Held-out vs modeling", "90", "0.598", "0.545", "0.441", "0.895", "0.303", "0.797", "0.455"],
          ])
    P(doc, "Table 9. External validation with the frozen AllStrains models.", italic=True, size=10)
    fig(doc, PLOTS / "external_01_pred_vs_exp.png",
        "Figure 51. External validation — predicted vs experimental pMIC (n = 92 unique molecules).")

    H(doc, "3.8 Sensitivity analysis (frozen main model)", 2)
    P(doc,
      "H37Rv unique molecules (n = 15,661) are entirely contained in AllStrainsExceptResistant. "
      "Full-set H37Rv metrics are therefore optimistic. The chemically honest H37Rv number is "
      "held-out versus Train (molecules that went to the Butina test split): n = 2,991, "
      "R2 = 0.437, RMSE = 0.699, AUC = 0.808, AP = 0.482.")
    P(doc,
      "Resistant unique molecules (n = 2,395) overlap the modeling set at 85.1%. The 356 "
      "molecules absent from modeling are the true sensitivity subset: R2 = 0.432, "
      "RMSE = 0.825, AUC = 0.913, AP = 0.681.")
    table(doc,
          ["Dataset", "Subset", "n", "n active", "R2", "RMSE", "ROC-AUC", "AP", "F1"],
          [
              ["H37Rv", "all (= overlap modeling)", "15,661", "2,476", "0.843", "0.390", "0.918", "0.816", "0.746"],
              ["H37Rv", "overlap Train", "12,670", "2,169", "0.926", "0.270", "0.932", "0.851", "0.783"],
              ["H37Rv", "held-out vs Train", "2,991", "307", "0.437", "0.699", "0.808", "0.482", "0.476"],
              ["Resistant", "all", "2,395", "535", "0.530", "0.651", "0.817", "0.591", "0.585"],
              ["Resistant", "overlap modeling", "2,039", "472", "0.546", "0.616", "0.798", "0.578", "0.583"],
              ["Resistant", "held-out vs modeling", "356", "63", "0.432", "0.825", "0.913", "0.681", "0.597"],
              ["Resistant", "held-out vs Train", "724", "118", "0.373", "0.785", "0.830", "0.552", "0.473"],
          ])
    P(doc, "Table 10. Sensitivity analysis with the frozen AllStrains models (threshold 0.050).",
      italic=True, size=10)
    fig(doc, PLOTS / "sens_H37Rv_pred_vs_exp.png",
        "Figure 52. Sensitivity — H37Rv predicted vs experimental pMIC (orange = overlap with modeling).")
    fig(doc, PLOTS / "sens_Resistant_pred_vs_exp.png",
        "Figure 53. Sensitivity — Resistant predicted vs experimental pMIC "
        "(blue = held-out vs modeling, n = 356).")

    H(doc, "3.9 Library screens", 2)
    table(doc,
          ["Library", "Output", "Predicted", "Skipped", "Active at P >= 0.050"],
          [
              ["Approved drugs", "approved_pmic.xlsx", "14,605", "875", "833 (5.7%)"],
              ["4FDN natural products", "4FDN-Natural-products-dock_pmic.xlsx", "1,923", "5", "7 (0.4%)"],
              ["CyanoMetDB V03 2024", "CyanoMetDB_V03_2024_pmic.xlsx", "3,059", "184", "347 (11.3%)"],
          ])
    P(doc, "Table 11. Virtual screening with the frozen AllStrains models. "
           "Active_predicted uses the OOF threshold 0.050. Rank by Active_probability / "
           "pMIC_predicted for prioritization. COCONUT was not rescreened in this run.",
      italic=True, size=10)

    # ── 4 Discussion ────────────────────────────────────────────────────────
    H(doc, "4. Discussion and limitations", 1)
    bullets(doc, [
        "Exact-SMILES collapse was the dominant dedup step (10,147 rows on the main table). "
        "3,014 groups had conflicting experimental pMIC under different conditions; median "
        "aggregation avoids letting outlier assay rows dominate.",
        "CV Average Precision is the appropriate classifier selection metric here: the "
        "scientific goal is recovery of a 10-16% active class, not overall accuracy. "
        "LightGBM (CV AP 0.751) was preferred over ExtraTrees/Stacking despite their "
        "slightly higher Test AUC.",
        "SMOTE was deliberately omitted. Fingerprint-space interpolation is not a valid "
        "molecule. class_weight plus OOF threshold is the chemically safer remedy.",
        "The OOF threshold of 0.050 sits at the search floor. It should be interpreted "
        "together with ranking metrics; probability scores remain the screening currency.",
        "H37Rv sensitivity on the full table is not an independent test (100% overlap). "
        "Report held-out versus Train (R2 0.437, AUC 0.808) instead.",
        "Resistant held-out versus modeling (n = 356, AUC 0.913, AP 0.681) is the strongest "
        "evidence that the AllStrains model transfers to a distinct labeled collection.",
        "External validation remains favorable (AUC 0.90) with only two overlapping molecules. "
        "AP on External is lower (0.30) because only 7 of 92 compounds are active at cutoff 6.",
        "Train classification metrics near 1.0 indicate overfitting of hard labels; OOF and "
        "Test are the valid estimates.",
        "SHAP bit maps are correlative attributions, not causal proofs of a pharmacophore. "
        "Sign is molecule-specific: morgan_1088 can be negative on one structure and positive on another.",
        "Cascade active-only regression was not retrained on this AllStrains split.",
        "COCONUT and full-library AD-coverage bars were not recomputed in this execution.",
    ])

    H(doc, "5. Key deliverables", 1)
    bullets(doc, [
        "cache_features_allstrains.joblib (features + Butina split)",
        "best_reg_model.joblib / best_cls_model.joblib (T = 6, decision_threshold = 0.050, selection_metric = cv_average_precision)",
        "preprocessed_data.xlsx (n = 18,918; Split, Dataset_role, Butina_cluster)",
        "plots/pmic_distribution.png and plots/cls_00c_oof_threshold.png",
        "plots/fp_highlights/ and plots/fp_highlights_cls/ (per-bit signed SHAP PNGs + CSVs)",
        "external_validation_predictions.xlsx / external_validation_metrics.csv",
        "sensitivity_H37Rv_predictions.xlsx / sensitivity_Resistant_predictions.xlsx / sensitivity_analysis_metrics.csv",
        "approved_pmic.xlsx, 4FDN-Natural-products-dock_pmic.xlsx, CyanoMetDB_V03_2024_pmic.xlsx",
        "pipeline_allstrains.log",
    ])

    H(doc, "6. Conclusions", 1)
    P(doc,
      "The pipeline was retrained on AllStrainsExceptResistant after two-stage SMILES "
      "deduplication with median pMIC (n = 18,918 unique molecules; 15.0% active at cutoff 6). "
      "Under a Butina split, stacking regression reached Test R2 = 0.41. Classification "
      "used CV Average Precision to select LightGBM (CV AP = 0.75; Test AUC = 0.78, Test AP = 0.44) "
      "and a frozen OOF threshold of 0.050. External validation on 92 unique molecules remained "
      "strong (R2 = 0.60; AUC = 0.90), with only two canonical overlaps. Sensitivity analysis "
      "showed that H37Rv is nested in the modeling set, while 356 Resistant molecules are true "
      "held-outs (AUC = 0.91). Approved, 4FDN, and CyanoMetDB were screened with the same "
      "frozen models. Per-bit SHAP maps now show individual fingerprint bits with signed "
      "positive or negative contribution.")

    end = doc.add_paragraph()
    end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = end.add_run(
        "\n- End of report -\n"
        "Generated from the AllStrainsExceptResistant execution "
        "(pipeline_allstrains.log, 25 September 2026)."
    )
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(100, 100, 100)

    doc.save(OUT)
    print(f"[saved] {OUT.resolve()}")
    return OUT


if __name__ == "__main__":
    build()
